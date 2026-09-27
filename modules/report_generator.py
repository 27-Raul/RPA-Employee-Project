"""
report_generator.py
===================
RPA Module: Onboarding Report Generator
Creates a formatted Excel report summarizing the onboarding batch.
"""

import os
import logging
from datetime import datetime

logger = logging.getLogger("RPA_Onboarding")


def generate_onboarding_report(results: list, report_folder: str,
                                company_name: str = "ABC Technologies") -> dict:
    """
    Generates an Excel onboarding report.

    Args:
        results: List of dicts, one per employee processed:
            {
              "employee_id":   "EMP001",
              "full_name":     "John Doe",
              "email":         "john@example.com",
              "department":    "IT",
              "designation":   "Developer",
              "status":        "Success" | "Failed",
              "folder_created": True/False,
              "letter_generated": True/False,
              "id_card_generated": True/False,
              "email_sent":    True/False,
              "db_updated":    True/False,
              "errors":        ["..."],
            }
        report_folder: Path where the report Excel file should be saved
        company_name:  From config

    Returns:
        {"success": True/False, "file_path": "...", "summary": {...}}
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side, numbers
        )

        os.makedirs(report_folder, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename  = f"OnboardingReport_{timestamp}.xlsx"
        file_path = os.path.join(report_folder, filename)

        wb = openpyxl.Workbook()

        # ── Summary Sheet ──────────────────────────────────────────────────────
        ws_sum = wb.active
        ws_sum.title = "Summary"

        total      = len(results)
        successful = sum(1 for r in results if r.get("status") == "Success")
        failed     = total - successful
        emails     = sum(1 for r in results if r.get("email_sent"))
        letters    = sum(1 for r in results if r.get("letter_generated"))
        id_cards   = sum(1 for r in results if r.get("id_card_generated"))
        folders    = sum(1 for r in results if r.get("folder_created"))

        dark_blue = "1E3A5F"
        gold      = "C9A84C"
        green     = "27AE60"
        red       = "E74C3C"
        light_bg  = "EEF3F8"

        def hfont(bold=False, color="000000", size=11):
            return Font(bold=bold, color=color, size=size)

        def hfill(hex_color):
            return PatternFill("solid", fgColor=hex_color)

        def hborder():
            s = Side(style="thin", color="CCCCCC")
            return Border(left=s, right=s, top=s, bottom=s)

        # Title
        ws_sum.merge_cells("A1:G1")
        ws_sum["A1"] = f"{company_name} — Employee Onboarding Report"
        ws_sum["A1"].font      = hfont(bold=True, color="FFFFFF", size=16)
        ws_sum["A1"].fill      = hfill(dark_blue)
        ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_sum.row_dimensions[1].height = 35

        # Generated date
        ws_sum.merge_cells("A2:G2")
        ws_sum["A2"] = f"Generated on: {datetime.now().strftime('%d %B %Y, %H:%M:%S')}"
        ws_sum["A2"].font      = hfont(color="666666", size=10)
        ws_sum["A2"].fill      = hfill("F5F5F5")
        ws_sum["A2"].alignment = Alignment(horizontal="center")
        ws_sum.row_dimensions[2].height = 18

        ws_sum.append([])  # empty row

        # KPI table
        kpi_rows = [
            ("Total Employees Processed", total,      dark_blue),
            ("Successful Onboardings",    successful,  green),
            ("Failed Onboardings",        failed,      red if failed else "999999"),
            ("Employee Folders Created",  folders,     "2980B9"),
            ("Welcome Letters Generated", letters,     "8E44AD"),
            ("ID Cards Generated",        id_cards,    "16A085"),
            ("Welcome Emails Sent",       emails,      "D35400"),
        ]

        for label, value, color in kpi_rows:
            row_idx = ws_sum.max_row + 1
            ws_sum.append([label, value])
            ws_sum[f"A{row_idx}"].font      = hfont(bold=True, color="333333", size=12)
            ws_sum[f"A{row_idx}"].fill      = hfill(light_bg)
            ws_sum[f"A{row_idx}"].border    = hborder()
            ws_sum[f"A{row_idx}"].alignment = Alignment(vertical="center")
            ws_sum[f"B{row_idx}"].font      = hfont(bold=True, color=color, size=14)
            ws_sum[f"B{row_idx}"].fill      = hfill("FFFFFF")
            ws_sum[f"B{row_idx}"].border    = hborder()
            ws_sum[f"B{row_idx}"].alignment = Alignment(horizontal="center", vertical="center")
            ws_sum.row_dimensions[row_idx].height = 22

        ws_sum.column_dimensions["A"].width = 35
        ws_sum.column_dimensions["B"].width = 20

        # ── Detailed Results Sheet ─────────────────────────────────────────────
        ws_det = wb.create_sheet("Detailed Results")

        headers = [
            "Employee ID", "Full Name", "Email", "Department", "Designation",
            "Status", "Folder Created", "Letter Generated", "ID Card", "Email Sent",
            "DB Updated", "Errors"
        ]
        ws_det.append(headers)

        # Header styling
        for cell in ws_det[1]:
            cell.font      = Font(bold=True, color="FFFFFF", size=10)
            cell.fill      = PatternFill("solid", fgColor=dark_blue)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = hborder()
        ws_det.row_dimensions[1].height = 20

        bool_icon = lambda b: "✅" if b else "❌"

        for i, r in enumerate(results):
            row = [
                r.get("employee_id", ""),
                r.get("full_name", ""),
                r.get("email", ""),
                r.get("department", ""),
                r.get("designation", ""),
                r.get("status", ""),
                bool_icon(r.get("folder_created")),
                bool_icon(r.get("letter_generated")),
                bool_icon(r.get("id_card_generated")),
                bool_icon(r.get("email_sent")),
                bool_icon(r.get("db_updated")),
                "; ".join(r.get("errors", [])) or "None",
            ]
            ws_det.append(row)
            row_idx = i + 2
            # Alternating row color
            bg = "FFFFFF" if i % 2 == 0 else light_bg
            status = r.get("status", "")
            for cell in ws_det[row_idx]:
                cell.fill   = PatternFill("solid", fgColor=bg)
                cell.border = hborder()
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            # Color the Status cell
            status_cell = ws_det.cell(row=row_idx, column=6)
            if status == "Success":
                status_cell.font = Font(bold=True, color=green)
            else:
                status_cell.font = Font(bold=True, color=red)

        # Auto column widths
        for col in ws_det.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws_det.column_dimensions[col[0].column_letter].width = min(max_len + 4, 35)
        ws_det.freeze_panes = "A2"

        wb.save(file_path)
        logger.info(f"Onboarding report saved: {file_path}")

        summary = {
            "total": total, "successful": successful, "failed": failed,
            "emails_sent": emails, "letters_generated": letters,
            "id_cards_generated": id_cards, "folders_created": folders,
        }
        return {"success": True, "file_path": file_path, "summary": summary}

    except Exception as e:
        msg = f"Failed to generate onboarding report: {e}"
        logger.error(msg)
        return {"success": False, "file_path": "", "summary": {}, "error": msg}
