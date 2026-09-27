"""
excel_handler.py
================
RPA Module: Excel Database Handler
Creates and updates EmployeeDatabase.xlsx and ErrorLog.xlsx.
"""

import os
import logging
from datetime import datetime

logger = logging.getLogger("RPA_Onboarding")

# ─── Column Definitions ────────────────────────────────────────────────────────
DB_HEADERS = [
    "Employee ID", "Full Name", "Date of Birth", "Gender", "Email",
    "Phone", "Address", "Department", "Designation", "Date of Joining",
    "Salary", "Manager Name", "Employment Type", "Status",
    "Onboarding Date", "Employee Folder Path", "Welcome Letter Path", "Email Status"
]

ERROR_HEADERS = [
    "Timestamp", "Employee ID", "Error Type", "Error Description", "Status"
]

# ─── Style helpers ─────────────────────────────────────────────────────────────
def _apply_header_style(ws):
    """Applies bold blue header style to row 1."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_font   = Font(bold=True, color="FFFFFF", size=11)
        header_fill   = PatternFill("solid", fgColor="1E3A5F")
        header_align  = Alignment(horizontal="center", vertical="center")
        thin_border   = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"),  bottom=Side(style="thin")
        )
        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align
            cell.border    = thin_border
        ws.row_dimensions[1].height = 20
    except Exception as e:
        logger.warning(f"Could not apply header style: {e}")


def _auto_column_widths(ws):
    """Auto-fits column widths based on content."""
    try:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)
    except Exception as e:
        logger.warning(f"Could not auto-fit columns: {e}")


# ─── Database Initialization ───────────────────────────────────────────────────
def initialize_database(database_path: str) -> bool:
    """
    Creates EmployeeDatabase.xlsx with headers if it does not exist.
    Returns True if successful.
    """
    try:
        import openpyxl
        if os.path.exists(database_path):
            logger.info("EmployeeDatabase.xlsx already exists. Skipping initialization.")
            return True

        os.makedirs(os.path.dirname(database_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employees"
        ws.append(DB_HEADERS)
        _apply_header_style(ws)
        _auto_column_widths(ws)
        wb.save(database_path)
        logger.info(f"EmployeeDatabase.xlsx created at: {database_path}")
        return True

    except Exception as e:
        logger.error(f"Failed to initialize database: {e}")
        return False


def initialize_error_log(error_log_path: str) -> bool:
    """Creates ErrorLog.xlsx with headers if it does not exist."""
    try:
        import openpyxl
        if os.path.exists(error_log_path):
            return True

        os.makedirs(os.path.dirname(error_log_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ErrorLog"
        ws.append(ERROR_HEADERS)
        _apply_header_style(ws)
        _auto_column_widths(ws)
        wb.save(error_log_path)
        logger.info(f"ErrorLog.xlsx created at: {error_log_path}")
        return True

    except Exception as e:
        logger.error(f"Failed to initialize error log: {e}")
        return False


# ─── Add Employee Record ───────────────────────────────────────────────────────
def add_employee_record(database_path: str, employee_data: dict) -> bool:
    """
    Adds a new employee row to EmployeeDatabase.xlsx with status = 'Onboarding'.

    Args:
        database_path: Full path to EmployeeDatabase.xlsx
        employee_data: Dict with all employee fields + 'employee_id'

    Returns:
        True if successful
    """
    try:
        import openpyxl
        initialize_database(database_path)
        wb = openpyxl.load_workbook(database_path)
        ws = wb.active

        row = [
            employee_data.get("employee_id", ""),
            employee_data.get("full_name", ""),
            employee_data.get("date_of_birth", ""),
            employee_data.get("gender", ""),
            employee_data.get("email", ""),
            employee_data.get("phone", ""),
            employee_data.get("address", ""),
            employee_data.get("department", ""),
            employee_data.get("designation", ""),
            employee_data.get("date_of_joining", ""),
            employee_data.get("salary", ""),
            employee_data.get("manager_name", ""),
            employee_data.get("employment_type", ""),
            "Onboarding",           # Status
            "",                     # Onboarding Date (filled later)
            "",                     # Employee Folder Path
            "",                     # Welcome Letter Path
            "Pending"               # Email Status
        ]
        ws.append(row)
        _auto_column_widths(ws)
        wb.save(database_path)
        logger.info(f"Employee record added: {employee_data.get('employee_id')}")
        return True

    except Exception as e:
        logger.error(f"Failed to add employee record: {e}")
        return False


# ─── Update Employee Record ────────────────────────────────────────────────────
def update_employee_record(database_path: str, employee_id: str, updates: dict) -> bool:
    """
    Updates an existing employee row in EmployeeDatabase.xlsx.

    Args:
        database_path: Full path to EmployeeDatabase.xlsx
        employee_id: The Employee ID to find (e.g. "EMP001")
        updates: Dict of {column_name: new_value}

    Returns:
        True if the row was found and updated
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(database_path)
        ws = wb.active

        # Build column name → column index map
        header_row = [cell.value for cell in ws[1]]
        col_map = {name: idx + 1 for idx, name in enumerate(header_row) if name}

        updated = False
        for row in ws.iter_rows(min_row=2):
            if row[0].value == employee_id:
                for field, value in updates.items():
                    if field in col_map:
                        ws.cell(row=row[0].row, column=col_map[field]).value = value
                updated = True
                break

        if updated:
            wb.save(database_path)
            logger.info(f"Employee record updated: {employee_id} → {list(updates.keys())}")
        else:
            logger.warning(f"Employee ID not found in database: {employee_id}")

        return updated

    except Exception as e:
        logger.error(f"Failed to update employee record ({employee_id}): {e}")
        return False


# ─── Log Error ─────────────────────────────────────────────────────────────────
def log_error(error_log_path: str, employee_id: str, error_type: str,
              error_description: str, status: str = "Failed") -> bool:
    """Appends an error entry to ErrorLog.xlsx."""
    try:
        import openpyxl
        initialize_error_log(error_log_path)
        wb = openpyxl.load_workbook(error_log_path)
        ws = wb.active

        ws.append([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            employee_id,
            error_type,
            error_description,
            status
        ])
        wb.save(error_log_path)
        return True
    except Exception as e:
        logger.error(f"Failed to write to ErrorLog.xlsx: {e}")
        return False


# ─── Read Config ───────────────────────────────────────────────────────────────
def read_config(config_path: str) -> dict:
    """
    Reads Config.xlsx (Key | Value columns) into a dictionary.

    Returns default values if file is missing or unreadable.
    """
    defaults = {
        "CompanyName":         "ABC Technologies",
        "HRName":              "HR Department",
        "HRDepartment":        "Human Resources",
        "EmployeeRootFolder":  "Employees",
        "DatabasePath":        "Database/EmployeeDatabase.xlsx",
        "ErrorLogPath":        "Logs/ErrorLog.xlsx",
        "ReportPath":          "Reports",
        "EmailMode":           "mock",
        "EmailSubject":        "Welcome to {CompanyName} - {EmployeeID}",
        "SMTPServer":          "smtp.gmail.com",
        "SMTPPort":            "587",
        "SenderEmail":         "",
        "SenderPassword":      "",
    }

    try:
        import openpyxl
        if not os.path.exists(config_path):
            logger.warning(f"Config.xlsx not found at {config_path}. Using defaults.")
            return defaults

        wb = openpyxl.load_workbook(config_path)
        ws = wb.active
        config = dict(defaults)  # start with defaults

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[1] is not None:
                config[str(row[0]).strip()] = str(row[1]).strip()

        return config

    except Exception as e:
        logger.error(f"Failed to read Config.xlsx: {e}")
        return defaults


# ─── Get All Employees ─────────────────────────────────────────────────────────
def get_all_employees(database_path: str) -> list:
    """Returns all employee records as a list of dicts."""
    try:
        import openpyxl
        if not os.path.exists(database_path):
            return []

        wb = openpyxl.load_workbook(database_path)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        employees = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                emp = {headers[i]: (row[i] if i < len(row) else "") for i in range(len(headers))}
                employees.append(emp)

        return employees
    except Exception as e:
        logger.error(f"Failed to read employees: {e}")
        return []


# ─── Excel Template Generator ───────────────────────────────────────────────────
TEMPLATE_HEADERS = [
    "Full Name", "Date of Birth (YYYY-MM-DD)", "Gender", "Email Address",
    "Phone Number", "Residential Address", "Department", "Designation",
    "Date of Joining (YYYY-MM-DD)", "Annual CTC (INR)", "Reporting Manager",
    "Employment Type", "Work Location", "Send Offer Email (Yes/No)"
]

SAMPLE_TEMPLATE_ROWS = [
    [
        "Vikram Sengupta", "1996-04-12", "Male", "vikram.sengupta@example.com",
        "9876512345", "42 Palm Avenue, Indiranagar, Bangalore", "IT", "Senior Systems Engineer",
        "2026-11-01", "1200000", "Deepak Mehta", "Full-Time", "Bangalore", "Yes"
    ],
    [
        "Ananya Deshmukh", "1998-09-25", "Female", "ananya.deshmukh@example.com",
        "9812345678", "15 Lakeview Residency, Hitec City, Hyderabad", "Engineering", "Software Developer",
        "2026-11-15", "950000", "Sunita Rao", "Full-Time", "Hyderabad", "Yes"
    ],
    [
        "Rahul Nambiar", "2000-01-18", "Male", "rahul.nambiar@example.com",
        "9745612389", "78 Green Glen Layout, Bellandur, Bangalore", "HR", "HR Executive",
        "2026-12-01", "650000", "Pooja Hegde", "Full-Time", "Bangalore", "Yes"
    ]
]


def create_employee_template(template_path: str) -> bool:
    """Creates a sample Employee Onboarding Excel template with guidance headers."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        os.makedirs(os.path.dirname(template_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employee_Onboarding_Template"

        # Headers
        ws.append(TEMPLATE_HEADERS)

        # Style header row
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="007CC3")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Sample rows
        for row in SAMPLE_TEMPLATE_ROWS:
            ws.append(row)

        for row_cells in ws.iter_rows(min_row=2, max_row=len(SAMPLE_TEMPLATE_ROWS) + 1):
            for cell in row_cells:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # Auto column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

        wb.save(template_path)
        logger.info(f"Excel onboarding template created at {template_path}")
        return True
    except Exception as e:
        logger.error(f"Failed to generate Excel template: {e}")
        return False


# ─── Parse Uploaded Employee Excel ──────────────────────────────────────────────
def parse_employee_excel(file_path_or_stream) -> list:
    """
    Parses an uploaded Excel file or BytesIO stream and extracts standardized
    employee candidate dictionaries for batch RPA onboarding.
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path_or_stream, data_only=True)
        ws = wb.active

        raw_headers = [str(cell.value or '').strip().lower() for cell in ws[1]]
        if not any(raw_headers):
            return []

        # Canonical key map mapping various possible header labels to internal fields
        FIELD_MAPPINGS = {
            "full_name": ["full name", "fullname", "name", "employee name", "candidate name", "first name"],
            "date_of_birth": ["date of birth", "dob", "birth date", "birthdate", "date of birth (yyyy-mm-dd)"],
            "gender": ["gender", "sex"],
            "email": ["email", "email address", "personal email", "candidate email", "mail"],
            "phone": ["phone", "phone number", "contact", "mobile", "mobile number", "contact phone"],
            "address": ["address", "residential address", "permanent address", "location address"],
            "department": ["department", "dept", "business unit"],
            "designation": ["designation", "role", "job role", "position", "job title", "title"],
            "date_of_joining": ["date of joining", "doj", "joining date", "joining date (yyyy-mm-dd)", "start date"],
            "salary": ["salary", "annual ctc", "annual ctc (inr)", "ctc", "compensation", "annual salary"],
            "manager_name": ["reporting manager", "manager", "manager name", "reporting to"],
            "employment_type": ["employment type", "emp type", "type", "job type"],
            "location": ["work location", "location", "base location", "city", "center"],
            "send_offer": ["send offer", "send offer email", "send offer email (yes/no)", "offer email", "send_offer"]
        }

        col_index_map = {}
        for canonical_field, aliases in FIELD_MAPPINGS.items():
            for idx, raw_h in enumerate(raw_headers):
                clean_h = raw_h.replace("_", " ").replace("-", " ").strip()
                if any(alias in clean_h or clean_h in alias for alias in aliases):
                    if canonical_field not in col_index_map:
                        col_index_map[canonical_field] = idx
                        break

        candidates = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue

            def get_val(field, default=""):
                idx = col_index_map.get(field)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    val = row[idx]
                    import datetime as dt_mod
                    if isinstance(val, (dt_mod.datetime, dt_mod.date)):
                        return val.strftime("%Y-%m-%d")
                    return str(val).strip()
                return default

            name = get_val("full_name")
            email = get_val("email")
            if not name and not email:
                continue

            # Format Dates cleanly
            dob = get_val("date_of_birth", "1998-01-01")
            if len(dob) > 10 and " " in dob:
                dob = dob.split()[0]

            doj = get_val("date_of_joining", "")
            if len(doj) > 10 and " " in doj:
                doj = doj.split()[0]
            if not doj:
                # Default DOJ to +14 days
                next_date = datetime.now()
                import datetime as dt_mod
                next_date += dt_mod.timedelta(days=14)
                doj = next_date.strftime("%Y-%m-%d")

            send_offer_val = get_val("send_offer", "yes").lower()
            send_offer_bool = send_offer_val in ["yes", "true", "1", "y", "enable"]

            candidate = {
                "row_number": row_idx,
                "full_name": name,
                "date_of_birth": dob,
                "gender": get_val("gender", "Female"),
                "email": email,
                "phone": get_val("phone", "9876543210"),
                "address": get_val("address", "Bangalore Center"),
                "department": get_val("department", "IT"),
                "designation": get_val("designation", "Systems Engineer"),
                "salary": str(get_val("salary", "850000")).replace(",", "").replace(".0", ""),
                "date_of_joining": doj,
                "manager_name": get_val("manager_name", "HR Manager"),
                "employment_type": get_val("employment_type", "Full-Time"),
                "location": get_val("location", "Bangalore"),
                "send_offer": send_offer_bool
            }
            candidates.append(candidate)

        return candidates
    except Exception as e:
        logger.error(f"Failed to parse employee Excel: {e}")
        return []

