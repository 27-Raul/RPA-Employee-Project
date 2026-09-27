import os
import random
import string
import logging
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

HEADERS = ['Asset ID', 'Asset Type', 'Serial Number', 'Employee ID', 'Employee Name', 'Assigned Date', 'Status', 'Notes']

def _style_header(ws):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="007CC3", end_color="007CC3", fill_type="solid")
    for col_num, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

def initialize_asset_db(path):
    """Initialize the AssetTracker Excel database."""
    try:
        if not os.path.exists(path):
            os.makedirs(os.path.dirname(path), exist_ok=True)
            wb = Workbook()
            ws = wb.active
            ws.title = "Assets"
            _style_header(ws)
            wb.save(path)
            logger.info(f"Initialized new Asset DB at {path}")
            return True
        logger.info(f"Asset DB already exists at {path}")
        return True
    except Exception as e:
        logger.error(f"Error initializing Asset DB: {e}")
        return False

def _generate_asset_id(path):
    try:
        wb = load_workbook(path)
        ws = wb.active
        max_row = ws.max_row
        if max_row <= 1:
            return "AST-0001"
        
        last_id = ws.cell(row=max_row, column=1).value
        if last_id and last_id.startswith("AST-"):
            num = int(last_id.split("-")[1])
            return f"AST-{num + 1:04d}"
        return "AST-0001"
    except Exception as e:
        logger.error(f"Error generating Asset ID: {e}")
        return "AST-0001"

def allocate_asset(path, employee_id, employee_name, asset_type, serial_number='', notes=''):
    """Allocate a single asset to an employee."""
    try:
        initialize_asset_db(path)
        asset_id = _generate_asset_id(path)
        assigned_date = datetime.now().strftime('%Y-%m-%d')
        status = 'Allocated'
        
        wb = load_workbook(path)
        ws = wb.active
        ws.append([asset_id, asset_type, serial_number, employee_id, employee_name, assigned_date, status, notes])
        wb.save(path)
        logger.info(f"Allocated asset {asset_id} to {employee_id}")
        return {'success': True, 'asset_id': asset_id, 'error': None}
    except Exception as e:
        logger.error(f"Error allocating asset: {e}")
        return {'success': False, 'asset_id': None, 'error': str(e)}

def get_employee_assets(path, employee_id):
    """Get all assets assigned to an employee."""
    assets = []
    try:
        if not os.path.exists(path):
            return assets
            
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[3] == employee_id:
                assets.append({
                    'asset_id': row[0],
                    'asset_type': row[1],
                    'serial_number': row[2],
                    'employee_id': row[3],
                    'employee_name': row[4],
                    'assigned_date': row[5],
                    'status': row[6],
                    'notes': row[7]
                })
        return assets
    except Exception as e:
        logger.error(f"Error retrieving employee assets: {e}")
        return assets

def get_all_assets(path):
    """Get all assets."""
    assets = []
    try:
        if not os.path.exists(path):
            return assets
            
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]: # Check if asset ID exists
                assets.append({
                    'asset_id': row[0],
                    'asset_type': row[1],
                    'serial_number': row[2],
                    'employee_id': row[3],
                    'employee_name': row[4],
                    'assigned_date': row[5],
                    'status': row[6],
                    'notes': row[7]
                })
        return assets
    except Exception as e:
        logger.error(f"Error retrieving all assets: {e}")
        return assets

def _generate_serial(asset_type):
    rand_str = ''.join(random.choices(string.ascii_uppercase + string.digits, k=4))
    if asset_type == 'Laptop':
        return f"LAP-2026-{rand_str}"
    elif asset_type == 'Access Card':
        return f"AC-{rand_str}"
    elif asset_type == 'Welcome Kit':
        return f"WK-{rand_str}"
    return f"SN-{rand_str}"

def auto_allocate_standard_kit(path, employee_id, employee_name):
    """Auto-allocate standard kit (Laptop, Access Card, Welcome Kit)."""
    results = {}
    standard_items = ['Laptop', 'Access Card', 'Welcome Kit']
    
    for item in standard_items:
        serial = _generate_serial(item)
        res = allocate_asset(path, employee_id, employee_name, item, serial, f"Auto-allocated standard {item}")
        results[item] = res
        
    return results

def update_asset_status(path, asset_id, new_status):
    """Update the status of an asset."""
    try:
        if not os.path.exists(path):
            logger.error(f"Database {path} does not exist.")
            return False
            
        wb = load_workbook(path)
        ws = wb.active
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == asset_id:
                ws.cell(row=row, column=7).value = new_status
                wb.save(path)
                logger.info(f"Updated status of {asset_id} to {new_status}")
                return True
                
        logger.warning(f"Asset ID {asset_id} not found.")
        return False
    except Exception as e:
        logger.error(f"Error updating asset status: {e}")
        return False
