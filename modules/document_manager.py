import os
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

logger = logging.getLogger(__name__)

REQUIRED_DOCUMENTS = [
    'Aadhaar Card',
    'PAN Card',
    'Passport Photo',
    '10th Marksheet',
    '12th Marksheet',
    'Degree Certificate',
    'Previous Experience Letter',
    'Bank Passbook / Cancelled Cheque'
]

HEADERS = [
    "Document ID", "Employee ID", "Employee Name", "Document Type",
    "File Name", "File Path", "Upload Date", "Verified",
    "Verified By", "Verified Date", "Notes"
]

def initialize_doc_tracker(path: str) -> bool:
    """Initialize the document tracking Excel file with headers."""
    try:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        if not os.path.exists(path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Document Tracker"

            header_fill = PatternFill(start_color="007CC3", end_color="007CC3", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)

            ws.append(HEADERS)
            for col_idx, _ in enumerate(HEADERS, 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
                ws.column_dimensions[col_letter].width = 18
            ws.column_dimensions['F'].width = 40
            
            wb.save(path)
            logger.info(f"Initialized Document Tracker at {path}")
        return True
    except Exception as e:
        logger.error(f"Failed to initialize document tracker: {e}")
        return False

def save_uploaded_document(upload_folder: str, employee_id: str, doc_type: str, file_obj, original_filename: str) -> dict:
    """
    Saves the uploaded file to: {upload_folder}/{emp_folder}/Personal_Documents/{doc_type}/{filename}
    Returns: {'success': bool, 'file_path': '...', 'error': '...'}
    """
    try:
        # Find matching employee folder in upload_folder (e.g. EMP001_Priya_Sharma or EMP001)
        target_emp_folder = None
        if os.path.exists(upload_folder):
            for item in os.listdir(upload_folder):
                item_path = os.path.join(upload_folder, item)
                if os.path.isdir(item_path) and (item == employee_id or item.startswith(f"{employee_id}_")):
                    target_emp_folder = item_path
                    break

        if not target_emp_folder:
            target_emp_folder = os.path.join(upload_folder, employee_id)

        target_dir = os.path.join(target_emp_folder, 'Personal_Documents', doc_type.replace("/", "_").replace("\\", "_"))
        os.makedirs(target_dir, exist_ok=True)

        safe_filename = original_filename.replace(" ", "_")
        file_path = os.path.join(target_dir, safe_filename)

        if hasattr(file_obj, 'save'):
            file_obj.save(file_path)
        else:
            with open(file_path, 'wb') as f:
                f.write(file_obj.read())

        logger.info(f"Saved document {doc_type} for {employee_id} at {file_path}")
        return {'success': True, 'file_path': file_path, 'error': None}
    except Exception as e:
        logger.error(f"Error saving document {doc_type} for {employee_id}: {e}")
        return {'success': False, 'file_path': None, 'error': str(e)}

def _generate_doc_id(ws) -> str:
    max_id = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        doc_id = row[0]
        if doc_id and isinstance(doc_id, str) and doc_id.startswith('DOC-'):
            try:
                num = int(doc_id.split('-')[1])
                if num > max_id:
                    max_id = num
            except ValueError:
                pass
    return f"DOC-{max_id + 1:04d}"

def add_document_record(tracker_path: str, employee_id: str, employee_name: str, doc_type: str, filename: str, file_path: str) -> bool:
    try:
        if not os.path.exists(tracker_path):
            initialize_doc_tracker(tracker_path)
            
        wb = openpyxl.load_workbook(tracker_path)
        ws = wb.active

        doc_id = _generate_doc_id(ws)
        upload_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        row = [
            doc_id, employee_id, employee_name, doc_type,
            filename, file_path, upload_date, "No",
            "", "", ""
        ]
        ws.append(row)
        wb.save(tracker_path)
        logger.info(f"Added document record {doc_id} for {employee_id}")
        return True
    except Exception as e:
        logger.error(f"Error adding document record for {employee_id}: {e}")
        return False

def get_employee_documents(tracker_path: str, employee_id: str) -> list:
    try:
        if not os.path.exists(tracker_path):
            return []
            
        wb = openpyxl.load_workbook(tracker_path, data_only=True)
        ws = wb.active
        
        docs = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] == employee_id:
                docs.append({
                    'document_id': row[0],
                    'employee_id': row[1],
                    'employee_name': row[2],
                    'doc_type': row[3],
                    'file_name': row[4],
                    'file_path': row[5],
                    'upload_date': row[6],
                    'verified': row[7] in ['Yes', True],
                    'verified_by': row[8],
                    'verified_date': row[9],
                    'notes': row[10]
                })
        return docs
    except Exception as e:
        logger.error(f"Error getting documents for {employee_id}: {e}")
        return []

def get_document_checklist(tracker_path: str, employee_id: str) -> list:
    """
    Returns a list of dicts showing all required docs with their upload and verification status.
    """
    try:
        uploaded_docs = get_employee_documents(tracker_path, employee_id)
        
        checklist = []
        for doc_type in REQUIRED_DOCUMENTS:
            doc_record = next((d for d in uploaded_docs if d['doc_type'] == doc_type), None)
            
            if doc_record:
                checklist.append({
                    'doc_type': doc_type,
                    'uploaded': True,
                    'verified': doc_record['verified'],
                    'file_path': doc_record['file_path'],
                    'document_id': doc_record['document_id']
                })
            else:
                checklist.append({
                    'doc_type': doc_type,
                    'uploaded': False,
                    'verified': False,
                    'file_path': None,
                    'document_id': None
                })
        return checklist
    except Exception as e:
        logger.error(f"Error generating document checklist for {employee_id}: {e}")
        return []

def verify_document(tracker_path: str, document_id: str, verified_by: str) -> bool:
    """Marks a document as verified with timestamp and verifier name."""
    try:
        if not os.path.exists(tracker_path):
            return False
            
        wb = openpyxl.load_workbook(tracker_path)
        ws = wb.active
        
        found = False
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if row[0] == document_id:
                ws.cell(row=row_idx, column=8, value="Yes")
                ws.cell(row=row_idx, column=9, value=verified_by)
                ws.cell(row=row_idx, column=10, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                found = True
                break
                
        if found:
            wb.save(tracker_path)
            logger.info(f"Verified document {document_id} by {verified_by}")
            return True
        return False
    except Exception as e:
        logger.error(f"Error verifying document {document_id}: {e}")
        return False

def get_verification_summary(tracker_path: str, employee_id: str) -> dict:
    try:
        checklist = get_document_checklist(tracker_path, employee_id)
        
        total = len(REQUIRED_DOCUMENTS)
        uploaded = sum(1 for d in checklist if d['uploaded'])
        verified = sum(1 for d in checklist if d['verified'])
        pending = total - uploaded
        
        percentage = (uploaded / total * 100) if total > 0 else 0
        
        return {
            'total_required': total,
            'uploaded': uploaded,
            'verified': verified,
            'pending': pending,
            'percentage': round(percentage, 2)
        }
    except Exception as e:
        logger.error(f"Error getting verification summary for {employee_id}: {e}")
        return {
            'total_required': len(REQUIRED_DOCUMENTS),
            'uploaded': 0,
            'verified': 0,
            'pending': len(REQUIRED_DOCUMENTS),
            'percentage': 0.0
        }
