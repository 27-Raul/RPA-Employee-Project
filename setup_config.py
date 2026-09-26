"""
setup_config.py (Upgraded for Infosys Enterprise Edition)
==========================================================
Run this ONCE to create Config.xlsx, initialize all databases, 
and seed sample employees.
Usage:  python setup_config.py
"""
import sys, io, os, json
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)


def create_config_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    path = os.path.join(BASE_DIR, "Config.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Config"

    config_data = [
        ("Key",                "Value",                              "Description"),
        ("CompanyName",        "Infosys Limited",                    "Your company name"),
        ("CompanyTagline",     "Navigate Your Next",                 "Company tagline"),
        ("HRName",             "HR Department",                      "HR contact name"),
        ("HRDepartment",       "Human Resources",                    "HR department name"),
        ("CompanyLocation",    "Bangalore, Karnataka, India",        "Primary office location"),
        ("EmployeeRootFolder", "Employees",                          "Root folder for employee folders"),
        ("DatabasePath",       "Database/EmployeeDatabase.xlsx",     "Path to HR database"),
        ("AssetDBPath",        "Database/AssetTracker.xlsx",         "Path to asset tracker"),
        ("DocTrackerPath",     "Database/DocumentTracker.xlsx",      "Path to document tracker"),
        ("ErrorLogPath",       "Logs/ErrorLog.xlsx",                 "Path to error log"),
        ("ReportPath",         "Reports",                            "Folder for onboarding reports"),
        ("EmailMode",          "mock",                               "mock = save to file, smtp = send"),
        ("EmailSubject",       "Welcome to {CompanyName} - {EmployeeID}", "Welcome email subject"),
        ("OfferEmailSubject",  "Offer of Employment - {CompanyName}","Offer letter email subject"),
        ("SMTPServer",         "smtp.gmail.com",                     "SMTP server"),
        ("SMTPPort",           "587",                                "SMTP port"),
        ("SenderEmail",        "",                                   "Gmail address (smtp mode)"),
        ("SenderPassword",     "",                                   "Gmail App Password (smtp mode)"),
        ("ProbationPeriod",    "6 months",                           "Employee probation period"),
        ("NoticePeriod",       "30 days",                            "Notice period"),
        ("WorkingHours",       "9:00 AM - 6:00 PM",                  "Standard working hours"),
    ]

    for i, row in enumerate(config_data):
        ws.append(row)
        r = ws.max_row
        if i == 0:
            for cell in ws[r]:
                cell.font      = Font(bold=True, color="FFFFFF", size=11)
                cell.fill      = PatternFill("solid", fgColor="007CC3")
                cell.alignment = Alignment(horizontal="center")
        else:
            ws.cell(r, 1).font = Font(bold=True, color="003366")
            ws.cell(r, 3).font = Font(italic=True, color="888888")

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 40
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"[OK] Config.xlsx created/updated: {path}")


def create_sample_employees():
    sample_employees = [
        {
            "full_name": "Priya Sharma", "date_of_birth": "1998-04-12",
            "gender": "Female", "email": "priya.sharma@example.com",
            "phone": "9876543210", "address": "42 MG Road, Bangalore - 560001",
            "department": "IT", "designation": "Software Engineer",
            "date_of_joining": "2026-10-01", "salary": "850000",
            "manager_name": "Rahul Kumar", "employment_type": "Full-Time",
            "location": "Bangalore"
        },
        {
            "full_name": "Arjun Mehta", "date_of_birth": "1995-08-22",
            "gender": "Male", "email": "arjun.mehta@example.com",
            "phone": "9123456789", "address": "17 Park Street, Mumbai - 400001",
            "department": "Finance", "designation": "Financial Analyst",
            "date_of_joining": "2026-10-15", "salary": "920000",
            "manager_name": "Sneha Patel", "employment_type": "Full-Time",
            "location": "Mumbai"
        },
        {
            "full_name": "Kavya Nair", "date_of_birth": "2000-01-30",
            "gender": "Female", "email": "kavya.nair@example.com",
            "phone": "8888123456", "address": "5 Anna Nagar, Chennai - 600040",
            "department": "HR", "designation": "HR Executive",
            "date_of_joining": "2026-11-01", "salary": "650000",
            "manager_name": "Deepa Krishnan", "employment_type": "Full-Time",
            "location": "Chennai"
        },
        {
            "full_name": "Rohan Verma", "date_of_birth": "1999-11-05",
            "gender": "Male", "email": "rohan.verma@example.com",
            "phone": "7777654321", "address": "88 Civil Lines, Delhi - 110054",
            "department": "Marketing", "designation": "Marketing Executive",
            "date_of_joining": "2026-10-20", "salary": "720000",
            "manager_name": "Neha Gupta", "employment_type": "Full-Time",
            "location": "Delhi"
        },
        {
            "full_name": "Sneha Pillai", "date_of_birth": "2001-06-18",
            "gender": "Female", "email": "sneha.pillai@example.com",
            "phone": "9900112233", "address": "33 Jubilee Hills, Hyderabad - 500033",
            "department": "R&D", "designation": "Research Intern",
            "date_of_joining": "2026-10-01", "salary": "300000",
            "manager_name": "Dr. Anil Menon", "employment_type": "Intern",
            "location": "Hyderabad"
        },
    ]

    input_dir = os.path.join(BASE_DIR, "Input")
    os.makedirs(input_dir, exist_ok=True)

    for i, emp in enumerate(sample_employees, 1):
        fname = f"Sample_Employee_{i:02d}_{emp['full_name'].replace(' ', '_')}.json"
        fpath = os.path.join(input_dir, fname)
        with open(fpath, "w") as f:
            json.dump(emp, f, indent=2)
        print(f"[OK] Sample: {fname}")


def create_demo_credentials_file():
    """Creates a DEMO_CREDENTIALS.txt for easy reference."""
    path = os.path.join(BASE_DIR, "DEMO_CREDENTIALS.txt")
    content = """
=======================================================
  Infosys HR Onboarding System — Demo Login Credentials
=======================================================

USERNAME      PASSWORD       ROLE          PERMISSIONS
------------------------------------------------------
admin         admin123       HR Admin      All features
hr            hr123          HR Staff      Register, upload, view
manager       manager123     Manager       View + approve
viewer        view123        Viewer        View only

URL: http://localhost:5000

Note: Salary column is hidden for Manager and Viewer roles.
=======================================================
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"[OK] Demo credentials file created: {path}")


if __name__ == "__main__":
    print("\n" + "="*60)
    print("  Infosys HR Onboarding System — Enterprise Setup")
    print("="*60)
    # Ensure all directories exist
    for d in ["Database", "Employees", "Reports", "Logs", "Input", "Templates"]:
        os.makedirs(os.path.join(BASE_DIR, d), exist_ok=True)

    create_config_xlsx()
    create_sample_employees()
    create_demo_credentials_file()

    # Initialize all Excel databases via their modules
    try:
        from modules.excel_handler import initialize_database, initialize_error_log
        from modules.asset_tracker import initialize_asset_db
        from modules.document_manager import initialize_doc_tracker

        initialize_database(os.path.join(BASE_DIR, "Database/EmployeeDatabase.xlsx"))
        initialize_error_log(os.path.join(BASE_DIR, "Logs/ErrorLog.xlsx"))
        initialize_asset_db(os.path.join(BASE_DIR, "Database/AssetTracker.xlsx"))
        initialize_doc_tracker(os.path.join(BASE_DIR, "Database/DocumentTracker.xlsx"))
        print("[OK] All databases initialized.")
    except Exception as e:
        print(f"[WARN] DB init error (run main.py to auto-init): {e}")

    print("\n[OK] Setup complete! Run:  python main.py")
    print("[OK] Then open:  http://localhost:5000")
    print("[OK] Login with: admin / admin123")
    print("="*60 + "\n")
